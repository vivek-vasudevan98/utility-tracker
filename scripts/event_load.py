"""Work out the daily Gateshead Suite load per event.

Reads the workbook written by extract_events.py and writes one row per
date and event name with a load score:

- Gateshead Suite (the whole suite) counts 3.
- A partition of it counts 1 per room it covers, so "Hillgate Suite" is 1
  and "Pipewellgate/Oakwellgate Suite" is 2.
- Gateshead Suite Foyer counts 1.

Bookings often split one package into several billable lines, so rows whose
event type contains "package" count once per date, event name and location.

Usage:
    python scripts/event_load.py report_events.xlsx [output.xlsx]
"""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FULL_SUITE = "gateshead suite"
FULL_SUITE_POINTS = 3
FOYER = "gateshead suite foyer"
FOYER_POINTS = 1
PARTITIONS = {"hillgate", "oakwellgate", "pipewellgate"}
PACKAGE_WORD = "package"

COLUMNS = ["Date", "Event Name", "Event Location", "Event Type"]


def _norm(value):
    """Lower-case and collapse whitespace (including non-breaking spaces)."""
    return " ".join(str(value).replace("\xa0", " ").split()).lower()


def room_points(location):
    """Load points for one booking of the given room, or None if unknown."""
    name = _norm(location)
    if name == FOYER:
        return FOYER_POINTS
    if name == FULL_SUITE:
        return FULL_SUITE_POINTS
    if name.endswith(" suite"):
        rooms = {p.strip() for p in name[: -len(" suite")].split("/")}
        if rooms and rooms <= PARTITIONS:
            return len(rooms)
    return None


def load_events(path):
    """Read the extract workbook, skipping its per-event group header rows."""
    df = pd.read_excel(path, dtype=object)
    missing = [c for c in COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"{path} is missing columns: {', '.join(missing)}")
    df = df[df["Event Name"].notna()].copy()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce", dayfirst=True)
    return df[df["Date"].notna()]


def calculate_load(events):
    events = events.copy()
    events["_is_package"] = events["Event Type"].map(
        lambda v: PACKAGE_WORD in _norm(v)
    )
    # Keep one package row per date, event name and location.
    keys = ["Date", "Event Name", "Event Location"]
    packages = events[events["_is_package"]].drop_duplicates(subset=keys)
    events = pd.concat([events[~events["_is_package"]], packages])

    events["Load"] = events["Event Location"].map(room_points)
    unknown = events.loc[events["Load"].isna(), "Event Location"].unique()
    if len(unknown):
        print(
            "Warning: no load rule for " + ", ".join(map(str, unknown))
            + "; counted as 0",
            file=sys.stderr,
        )
    events["Load"] = events["Load"].fillna(0).astype(int)

    load = events.groupby(["Date", "Event Name"], as_index=False)["Load"].sum()
    return load.sort_values(["Date", "Event Name"], kind="stable").reset_index(
        drop=True
    )


def write_load(df, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Event Load"
    columns = ["Date", "Event Name", "Load"]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for rec in df[columns].itertuples(index=False):
        ws.append([rec[0].to_pydatetime(), rec[1], int(rec[2])])
        ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
    for i, col in enumerate(columns, 1):
        width = max([len(col)] + [len(str(v)) for v in df[col]]) if len(df) else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 12) + 2, 60)
    wb.save(path)


def default_output(path):
    stem = path.stem
    if stem.endswith("_events"):
        stem = stem[: -len("_events")]
    return path.with_name(f"{stem}_eventload.xlsx")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input", type=Path, help="xlsx written by extract_events.py")
    parser.add_argument(
        "output",
        type=Path,
        nargs="?",
        help="output .xlsx file (default: <report>_eventload.xlsx)",
    )
    args = parser.parse_args(argv)
    output = args.output or default_output(args.input)

    load = calculate_load(load_events(args.input))
    write_load(load, output)
    print(f"Wrote {len(load)} date/event rows to {output}")


if __name__ == "__main__":
    sys.exit(main())
